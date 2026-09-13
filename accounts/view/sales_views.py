from decimal import Decimal, ROUND_HALF_UP
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db import transaction
from django.db.models import (
    Sum,
    F,
    Value,
    DecimalField,
    ExpressionWrapper,
    Max,
)
from django.http import JsonResponse, HttpResponseNotAllowed
from django.shortcuts import get_object_or_404, render
from django.utils import timezone
from django.views import View

# --- Adjust these imports to your project ---
from accounts.models import Product  
from accounts.models.stock import StockBatch, StockLedger  
from accounts.models.org import Branch                     
from accounts.models.sales import SalesHeader, SalesLine 
from accounts.models.Insurance import (
    Insurance,
    InsuranceScheme
)
from accounts.models.sale_reversal import SaleReversal      
from django.utils.dateparse import parse_date
from django.db.models import Q
from django.db.models import Case, When, Value, IntegerField
import csv
from django.http import HttpResponse
from openpyxl import Workbook
from io import BytesIO
from openpyxl import Workbook
from accounts.models.purchase import PurchaseHeader, PurchaseLine, Supplier
from django.db.models.functions import Coalesce, Cast, TruncDate
from django.http import JsonResponse
from django.views import View
from accounts.models.Insurance import InsuranceScheme
from accounts.utils.stock import average_purchase_cost


class InsuranceSchemeApiView(LoginRequiredMixin, View):
    def get(self, request):

        insurance_id = request.GET.get("insurance_id")

        if not insurance_id:
            return JsonResponse([], safe=False)

        schemes = (
            InsuranceScheme.objects
            .filter(
                insurance_id=insurance_id,
                is_active=True
            )
            .order_by("name")
        )

        data = [
            {
                "id": s.id,
                "name": s.name,
                "discount_percent": float(
                    s.discount_percent or 0
                )
            }
            for s in schemes
        ]

        return JsonResponse(data, safe=False)
    
class SaleDetailView(LoginRequiredMixin, View):
    """
    GET /sales/<uuid:sale_id>/
      - HTML receipt (default)
    GET /sales/<uuid:sale_id>/?format=json or Accept: application/json
      - JSON payload (for programmatic use)
    """
    def get(self, request, sale_id):
        header = (
            SalesHeader.objects
            .select_related("branch", "cashier")
            .prefetch_related("lines__product", "lines__batch")
            .filter(sale_id=sale_id)
            .first()
        )
        if not header:
            return render(request, "accounts/404.html", status=404)  # or raise 404

        # reversal info (there could be multiple in theory—grab the latest)
        reversal = None
        if hasattr(header, "reversal"):
            # related_name="reversal" gives a RelatedManager
            reversal = header.reversal.order_by("-created_at").first()

        # build useful aggregates
        lines = list(header.lines.all())

        subtotal = sum((l.unit_price * l.qty for l in lines), Decimal("0"))
        discount_total = sum((l.discount for l in lines), Decimal("0"))
        tax_total = header.tax_total or Decimal("0")
        grand_total = header.total or (subtotal - discount_total + tax_total)

        # JSON?
        wants_json = (
            request.GET.get("format") == "json"
            or "application/json" in (request.headers.get("Accept", "") or "")
        )
        if wants_json:
            payload = {
                "sale_id": str(header.sale_id),
                "status": header.status,
                "branch": getattr(header.branch, "name", None),
                "cashier": getattr(header.cashier, "username", None),
                "pos_terminal": header.pos_terminal,
                "created_at": timezone.localtime(header.created_at).isoformat(),
                "subtotal": float(subtotal),
                "discount_total": float(discount_total),
                "tax_total": float(tax_total),
                "total": float(grand_total),
                "payment_method": header.payment_method,
                "mpesa_mode": getattr(
                    header,
                    "mpesa_mode",
                    None,
                ),
                "payment_reference": getattr(
                    header,
                    "payment_reference",
                    None,
                ),
                "payment_phone": getattr(
                    header,
                    "payment_phone",
                    None,
                ),
                "insurance": (header.insurance.name if header.insurance_id else None),
                "insurance_scheme": (header.insurance_scheme.name if header.insurance_scheme_id else None),
                "member_no": header.member_no,
                "member_name": header.member_name,
                "auth_no": header.auth_no,
                "reversal": {
                    "reversed_by": getattr(reversal.reversed_by, "username", None) if reversal else None,
                    "reason": reversal.reason if reversal else None,
                    "created_at": timezone.localtime(reversal.created_at).isoformat() if reversal else None,
                },
                "lines": [{
                    "product_id": l.product_id,
                    "sku": getattr(l.product, "sku", None),
                    "name": getattr(l.product, "name", None),
                    "batch_no": getattr(l.batch, "batch_no", None) if l.batch_id else None,
                    "expiry_date": getattr(l.batch, "expiry_date", None) if l.batch_id else None,
                    "qty": float(l.qty),
                    "unit_price": float(l.unit_price),
                    "discount": float(l.discount or 0),
                    "line_total": float(l.line_total or (l.unit_price * l.qty - (l.discount or 0))),
                } for l in lines],
            }
            return JsonResponse(payload, status=200)

        # HTML context
        context = {
            "header": header,
            "lines": lines,
            "subtotal": subtotal,
            "discount_total": discount_total,
            "tax_total": tax_total,
            "grand_total": grand_total,
            "reversal": reversal,
        }
        return render(request, "accounts/sales_detail.html", context)

    def post(self, request, *_args, **_kwargs):
        return HttpResponseNotAllowed(["GET"])



# ---------- helpers ----------
def _q4(x):  # 4 dp quantization for qty, unit_price, discount, line_cost/total at line granularity
    return (Decimal(x or 0)).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)

def _q2(x):  # 2 dp totals on header
    return (Decimal(x or 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

def _active_branch(request):
    bid = request.session.get("active_branch_id")
    if not bid:
        return None
    return Branch.objects.filter(pk=bid).first()


def _average_cost(branch_id, product_id):
    return _q4(average_purchase_cost(product_id, branch_id))


# ---------- FEFO allocation (earliest expiry first; no-expiry last) ----------
def _fefo_batches(branch_id, product_id):
    """
    Returns unexpired batches with stock for (branch, product),
    ordered: non-null expiry ascending (earliest first), then null expiry.
    """
    today = timezone.localdate()
    return (
        StockBatch.objects
        .select_for_update()
        .filter(
            branch_id=branch_id,
            product_id=product_id,
            qty_on_hand__gt=0,
        )
        .filter(
            # allow no-expiry or expiry strictly greater than today
            # (block expired and expiring today)
            Q(expiry_date__isnull=True) | Q(expiry_date__gt=today)
        )
        .annotate(
            expiry_null=Case(
                When(expiry_date__isnull=True, then=Value(1)),
                default=Value(0),
                output_field=IntegerField(),
            )
        )
        .order_by("expiry_null", "expiry_date", "batch_no")
    )


def _ensure_available(branch_id, product_id, qty_needed: Decimal):
    today = timezone.localdate()
    total = (
        StockBatch.objects
        .filter(
            branch_id=branch_id,
            product_id=product_id,
        )
        .filter(
            Q(expiry_date__isnull=True) | Q(expiry_date__gt=today)
        )
        .aggregate(v=Sum('qty_on_hand'))['v'] or Decimal('0')
    )
    return Decimal(total) >= qty_needed


class SubmitSaleView(LoginRequiredMixin, View):
    """
    POST /sales/submit/

    Examples:

    Cash:
    {
        "pos_terminal": "WEB",
        "payment_method": "CASH",
        "lines": [...]
    }

    Card:
    {
        "pos_terminal": "WEB",
        "payment_method": "CARD",
        "payment_reference": "OPTIONAL-TERMINAL-REFERENCE",
        "lines": [...]
    }

    Manual M-Pesa:
    {
        "pos_terminal": "WEB",
        "payment_method": "MPESA",
        "mpesa_mode": "MANUAL",
        "payment_reference": "TQ12ABCDEF",
        "lines": [...]
    }

    Confirmed STK M-Pesa:
    {
        "pos_terminal": "WEB",
        "payment_method": "MPESA",
        "mpesa_mode": "STK",
        "payment_reference": "MPESA-RECEIPT-NUMBER",
        "phone": "2547XXXXXXXX",
        "lines": [...]
    }
    """

    VALID_PAYMENT_METHODS = {
        "CASH",
        "CARD",
        "MPESA",
        "INSURANCE",
    }

    VALID_MPESA_MODES = {
        "STK",
        "MANUAL",
    }

    def post(self, request, *args, **kwargs):
        import json

        # ========================================================
        # PARSE REQUEST
        # ========================================================

        try:
            body = json.loads(
                request.body.decode("utf-8")
            )
        except (json.JSONDecodeError, UnicodeDecodeError):
            return JsonResponse(
                {
                    "error": "Invalid JSON body."
                },
                status=400,
            )

        if not isinstance(body, dict):
            return JsonResponse(
                {
                    "error": "Request body must be a JSON object."
                },
                status=400,
            )

        lines = body.get("lines") or []

        if not isinstance(lines, list) or not lines:
            return JsonResponse(
                {
                    "error": "No sale lines provided."
                },
                status=400,
            )

        branch = _active_branch(request)

        if not branch:
            return JsonResponse(
                {
                    "error": "No active branch in session."
                },
                status=400,
            )

        cashier = request.user

        pos_terminal = str(
            body.get("pos_terminal") or "WEB"
        ).strip()[:20]

        # Important: normalize every payment value.
        payment_method = str(
            body.get("payment_method") or ""
        ).strip().upper()

        mpesa_mode = str(
            body.get("mpesa_mode") or ""
        ).strip().upper() or None

        payment_reference = str(
            body.get("payment_reference") or ""
        ).strip().upper() or None

        payment_phone = str(
            body.get("phone") or ""
        ).strip() or None

        if payment_method not in self.VALID_PAYMENT_METHODS:
            return JsonResponse(
                {
                    "error": (
                        "Invalid payment method. "
                        "Select CASH, CARD, MPESA or INSURANCE."
                    )
                },
                status=400,
            )

        # ========================================================
        # PAYMENT-SPECIFIC VALIDATION
        # ========================================================

        insurance = None
        scheme = None

        member_no = None
        member_name = None
        auth_no = None

        if payment_method == "MPESA":
            if mpesa_mode not in self.VALID_MPESA_MODES:
                return JsonResponse(
                    {
                        "error": (
                            "Select an M-Pesa payment mode: "
                            "STK or MANUAL."
                        )
                    },
                    status=400,
                )

            if mpesa_mode == "MANUAL":
                if not payment_reference:
                    return JsonResponse(
                        {
                            "error": (
                                "M-Pesa transaction code is "
                                "required for manual payment."
                            )
                        },
                        status=400,
                    )

                if not (
                    8 <= len(payment_reference) <= 20
                    and payment_reference.isalnum()
                ):
                    return JsonResponse(
                        {
                            "error": (
                                "Enter a valid M-Pesa "
                                "transaction code."
                            )
                        },
                        status=400,
                    )

                # Manual M-Pesa does not require a phone number.
                payment_phone = None

            elif mpesa_mode == "STK":
                if not payment_reference:
                    return JsonResponse(
                        {
                            "error": (
                                "The STK payment has no M-Pesa "
                                "receipt reference. Confirm the "
                                "payment before saving the sale."
                            )
                        },
                        status=400,
                    )

        else:
            # Prevent Card, Cash or Insurance from retaining
            # an old M-Pesa mode or phone number.
            mpesa_mode = None
            payment_phone = None

        if payment_method == "INSURANCE":
            insurance_id = body.get("insurance_id")
            scheme_id = body.get(
                "insurance_scheme_id"
            )

            if not insurance_id:
                return JsonResponse(
                    {
                        "error": "Insurance is required."
                    },
                    status=400,
                )

            if not scheme_id:
                return JsonResponse(
                    {
                        "error": (
                            "Insurance scheme is required."
                        )
                    },
                    status=400,
                )

            insurance = (
                Insurance.objects
                .filter(
                    pk=insurance_id,
                    is_active=True,
                )
                .first()
            )

            if not insurance:
                return JsonResponse(
                    {
                        "error": "Insurance not found."
                    },
                    status=404,
                )

            scheme = (
                InsuranceScheme.objects
                .filter(
                    pk=scheme_id,
                    insurance=insurance,
                    is_active=True,
                )
                .first()
            )

            if not scheme:
                return JsonResponse(
                    {
                        "error": (
                            "Insurance scheme not found."
                        )
                    },
                    status=404,
                )

            member_no = str(
                body.get("member_no") or ""
            ).strip() or None

            member_name = str(
                body.get("member_name") or ""
            ).strip() or None

            auth_no = str(
                body.get("auth_no") or ""
            ).strip() or None

            # Insurance does not use a normal payment reference.
            payment_reference = None

        else:
            # Prevent non-insurance sales from retaining
            # insurance information.
            insurance = None
            scheme = None
            member_no = None
            member_name = None
            auth_no = None

        # ========================================================
        # VALIDATE SALE LINES AND PRICES
        # ========================================================

        product_cache = {}
        validated_lines = []

        for index, line in enumerate(lines, start=1):
            if not isinstance(line, dict):
                return JsonResponse(
                    {
                        "error": (
                            f"Line {index}: invalid line data."
                        )
                    },
                    status=400,
                )

            product_id = line.get("product_id")

            if not product_id:
                return JsonResponse(
                    {
                        "error": (
                            f"Line {index}: "
                            "product_id is required."
                        )
                    },
                    status=400,
                )

            try:
                quantity = _q4(
                    line.get("qty")
                )

                entered_price = _q4(
                    line.get("unit_price")
                )

                entered_discount = _q4(
                    line.get("discount")
                )
            except Exception:
                return JsonResponse(
                    {
                        "error": (
                            f"Line {index}: invalid numeric value."
                        )
                    },
                    status=400,
                )

            if quantity <= 0:
                return JsonResponse(
                    {
                        "error": (
                            f"Line {index}: qty must be greater "
                            "than zero."
                        )
                    },
                    status=400,
                )

            if entered_price < 0:
                return JsonResponse(
                    {
                        "error": (
                            f"Line {index}: unit price cannot "
                            "be negative."
                        )
                    },
                    status=400,
                )

            if entered_discount < 0:
                return JsonResponse(
                    {
                        "error": (
                            f"Line {index}: discount cannot "
                            "be negative."
                        )
                    },
                    status=400,
                )

            if product_id not in product_cache:
                product = (
                    Product.objects
                    .filter(
                        pk=product_id,
                        active=True,
                    )
                    .first()
                )

                if not product:
                    return JsonResponse(
                        {
                            "error": (
                                f"Line {index}: product "
                                f"{product_id} was not found "
                                "or is inactive."
                            )
                        },
                        status=404,
                    )

                product_cache[product_id] = product

            product = product_cache[product_id]

            if payment_method == "INSURANCE":
                # Existing insurance calculation retained.
                average_cost = _average_cost(branch.id, product.id)
                unit_price = _q4(
                    average_cost
                    * Decimal(scheme.discount_percent)
                    / Decimal("100")
                )

                discount = Decimal("0")

            else:
                average_cost = _average_cost(branch.id, product.id)
                minimum_price = _q4(
                    average_cost
                    * Decimal("1.33")
                )

                maximum_price = _q4(
                    average_cost
                    * Decimal("1.50")
                )

                if average_cost <= 0:
                    return JsonResponse(
                        {"error": f"{product.name}: no purchase cost is available for this branch."},
                        status=400,
                    )

                if entered_price < minimum_price:
                    return JsonResponse(
                        {
                            "error": (
                                f"{product.name}: selling price "
                                "is below the minimum allowed "
                                f"price ({minimum_price})."
                            )
                        },
                        status=400,
                    )

                if entered_price > maximum_price:
                    return JsonResponse(
                        {
                            "error": (
                                f"{product.name}: selling price "
                                "is above the maximum allowed "
                                f"price ({maximum_price})."
                            )
                        },
                        status=400,
                    )

                unit_price = entered_price
                discount = entered_discount

            line_subtotal = _q4(
                unit_price * quantity
            )

            if discount > line_subtotal:
                return JsonResponse(
                    {
                        "error": (
                            f"{product.name}: discount cannot "
                            "exceed the line subtotal."
                        )
                    },
                    status=400,
                )

            validated_lines.append(
                {
                    "index": index,
                    "product_id": product_id,
                    "product": product,
                    "qty": quantity,
                    "unit_price": unit_price,
                    "discount": discount,
                }
            )

        # ========================================================
        # CREATE SALE AND ALLOCATE STOCK
        # ========================================================

        try:
            with transaction.atomic():
                # Lock and check stock inside the same transaction.
                batch_cache = {}

                for line in validated_lines:
                    product_id = line["product_id"]
                    quantity = line["qty"]

                    batches = list(
                        _fefo_batches(
                            branch.id,
                            product_id,
                        )
                    )

                    total_available = sum(
                        (
                            _q4(batch.qty_on_hand)
                            for batch in batches
                        ),
                        Decimal("0"),
                    )

                    if total_available < quantity:
                        return JsonResponse(
                            {
                                "error": (
                                    f"Line {line['index']}: "
                                    "insufficient stock for "
                                    f"{line['product'].name}."
                                )
                            },
                            status=400,
                        )

                    batch_cache[product_id] = batches

                # Only create the header after all validation
                # and initial stock checks have passed.
                header_fields = {
                    "branch": branch,
                    "cashier": cashier,
                    "pos_terminal": pos_terminal,
                    "payment_method": payment_method,
                    "mpesa_mode": mpesa_mode,
                    "insurance": insurance,
                    "insurance_scheme": scheme,
                    "member_no": member_no,
                    "member_name": member_name,
                    "auth_no": auth_no,
                    "total": _q2(0),
                    "tax_total": _q2(0),
                    "status": "POSTED",
                    "created_at": timezone.now(),
                }

                

                # Save these fields only if they exist in your model.
                sales_header_field_names = {
                    field.name
                    for field in SalesHeader._meta.fields
                }

                if (
                    "payment_reference"
                    in sales_header_field_names
                ):
                    header_fields[
                        "payment_reference"
                    ] = payment_reference

                if (
                    "payment_phone"
                    in sales_header_field_names
                ):
                    header_fields[
                        "payment_phone"
                    ] = payment_phone

                header = SalesHeader.objects.create(
                    **header_fields
                )

                total_ex = Decimal("0.00")
                tax_total = Decimal("0.00")

                for line in validated_lines:
                    product_id = line["product_id"]
                    quantity_to_allocate = line["qty"]
                    unit_price = line["unit_price"]
                    discount_remaining = line["discount"]

                    remaining = quantity_to_allocate

                    # Existing proportional discount
                    # calculation retained.
                    discount_per_unit = (
                        discount_remaining
                        / quantity_to_allocate
                        if quantity_to_allocate > 0
                        else Decimal("0")
                    )

                    batches = batch_cache[product_id]

                    for batch in batches:
                        if remaining <= 0:
                            break

                        take = min(
                            _q4(batch.qty_on_hand),
                            remaining,
                        )

                        if take <= 0:
                            continue

                        # Existing discount allocation retained.
                        portion_discount = _q4(
                            discount_per_unit * take
                        )

                        # Existing line total calculation retained.
                        line_total = _q4(
                            (unit_price * take)
                            - portion_discount
                        )

                        line_cost = _q4(batch.buying_cost)

                        updated = (
                            StockBatch.objects
                            .filter(
                                pk=batch.pk,
                                qty_on_hand__gte=take,
                            )
                            .update(
                                qty_on_hand=(
                                    F("qty_on_hand") - take
                                )
                            )
                        )

                        if updated == 0:
                            # Raising an exception guarantees
                            # transaction rollback.
                            raise RuntimeError(
                                "Stock changed while processing "
                                f"{line['product'].name}. "
                                "Please retry the sale."
                            )

                        SalesLine.objects.create(
                            header=header,
                            product_id=product_id,
                            batch=batch,
                            qty=take,
                            unit_price=unit_price,
                            discount=portion_discount,
                            line_total=line_total,
                            line_cost=line_cost,
                        )

                        StockLedger.objects.create(
                            product_id=product_id,
                            branch_id=branch.id,
                            batch_id=batch.pk,
                            qty_change=-take,
                            unit_cost=line_cost,
                            reason=StockLedger.OUT,
                            reference=str(
                                header.sale_id
                            ),
                        )

                        # Existing header total accumulation retained.
                        total_ex += _q2(line_total)
                        remaining -= take

                    if remaining > 0:
                        raise RuntimeError(
                            "Allocation logic left "
                            f"{remaining} units unallocated for "
                            f"{line['product'].name}."
                        )

                # Existing header total calculation retained.
                header.total = _q2(
                    total_ex + tax_total
                )

                header.tax_total = _q2(
                    tax_total
                )

                header.save(
                    update_fields=[
                        "total",
                        "tax_total",
                    ]
                )

                return JsonResponse(
                    {
                        "sale_id": str(
                            header.sale_id
                        ),
                        "total": float(
                            header.total
                        ),
                        "lines": header.lines.count(),
                        "status": header.status,

                        # Return the saved method so the browser
                        # can confirm what was actually recorded.
                        "payment_method":
                            header.payment_method,

                        "mpesa_mode":
                            getattr(
                                header,
                                "mpesa_mode",
                                None,
                            ),

                        "payment_reference":
                            getattr(
                                header,
                                "payment_reference",
                                None,
                            ),
                    },
                    status=200,
                )

        except RuntimeError as error:
            return JsonResponse(
                {
                    "error": str(error)
                },
                status=409,
            )

        except Exception as error:
            # Log the full error in your server console.
            import logging

            logger = logging.getLogger(__name__)

            logger.exception(
                "Failed to submit sale"
            )

            return JsonResponse(
                {
                    "error": (
                        "The sale could not be completed. "
                        f"{str(error)}"
                    )
                },
                status=500,
            )


class ReverseSaleView(LoginRequiredMixin, View):
    """
    POST /sales/reverse/
    Body: {"sale_id": "uuid-string", "reason": "optional text"}
    Creates a SaleReversal, restores stock to the same batches used in the sale (when present),
    sets header.status='REVERSED'.
    """
    def post(self, request, *args, **kwargs):
        import json
        try:
            body = json.loads(request.body.decode('utf-8'))
        except Exception:
            return JsonResponse({"error": "Invalid JSON body."}, status=400)

        sale_id = (body.get('sale_id') or '').strip()
        reason = (body.get('reason') or '').strip()[:300]
        mpesa_mode = body.get("mpesa_mode")

        if not sale_id:
            return JsonResponse({"error": "sale_id is required."}, status=400)

        header = SalesHeader.objects.filter(sale_id=sale_id).first()
        if not header:
            return JsonResponse({"error": "Sale not found."}, status=404)
        if header.status == "REVERSED":
            return JsonResponse({"error": "Sale already reversed."}, status=400)

        with transaction.atomic():
            # Create reversal record
            SaleReversal.objects.create(
                header=header,
                reversed_by=request.user,
                reason=reason or None,
            )

            # Restore batch qty_on_hand and write ledger (IN) for each original line
            for ln in header.lines.select_related('batch'):
                qty = _q4(ln.qty)
                if qty <= 0:
                    continue

                if ln.batch_id:
                    # Put back into the same batch
                    StockBatch.objects.filter(pk=ln.batch_id).update(qty_on_hand=F('qty_on_hand') + qty)

                    StockLedger.objects.create(
                        product_id=ln.product_id,
                        branch_id=header.branch_id,
                        batch_id=ln.batch_id,
                        qty_change=qty,
                        unit_cost=ln.batch.buying_cost,
                        reason=StockLedger.IN,
                        reference=str(header.sale_id),
                    )
                else:
                    # No batch recorded: use (create or reuse) a reversal batch marker
                    # Ensures unique combination per constraints
                    rev_batch_no = f"REV-{str(header.sale_id)[:8]}"
                    batch, _ = StockBatch.objects.get_or_create(
                        product_id=ln.product_id,
                        branch_id=header.branch_id,
                        batch_no=rev_batch_no,
                        expiry_date=None,
                        defaults={"qty_on_hand": Decimal('0'), "buying_cost": Decimal('0')}
                    )
                    StockBatch.objects.filter(pk=batch.pk).update(qty_on_hand=F('qty_on_hand') + qty)

                    StockLedger.objects.create(
                        product_id=ln.product_id,
                        branch_id=header.branch_id,
                        batch_id=batch.pk,
                        qty_change=qty,
                        unit_cost=batch.buying_cost,
                        reason=StockLedger.IN,
                        reference=str(header.sale_id),
                    )

            header.status = "REVERSED"
            header.save(update_fields=['status'])

            return JsonResponse({
                "sale_id": str(header.sale_id),
                "status": header.status,
                "reversed_at": timezone.now().isoformat(),
            }, status=200)



def _active_branch(request):
    bid = request.session.get("active_branch_id")
    return Branch.objects.filter(pk=bid).first() if bid else None

class SetActiveBranchView(LoginRequiredMixin, View):
    def post(self, request):
        branch_id = request.POST.get("branch_id")

        if not branch_id:
            return JsonResponse({"error": "branch_id required"}, status=400)

        link = request.user.branch_links.filter(branch_id=branch_id).first()
        if not link:
            return JsonResponse({"error": "Not assigned to this branch"}, status=403)

        request.session["active_branch_id"] = int(branch_id)
        return JsonResponse({"success": True})



class ReversalListApiView(LoginRequiredMixin, View):
    """
    GET /api/sales/reversals/?from=YYYY-MM-DD&to=YYYY-MM-DD&q=term
    Returns: list of dicts [{sale_id,total,reversed_by,reason,reversed_at},...]
    """
    def get(self, request, *args, **kwargs):
        br = _active_branch(request)
        qs = SaleReversal.objects.select_related('header', 'reversed_by')

        if br:
            qs = qs.filter(header__branch=br)

        # filters
        dfrom = parse_date(request.GET.get('from') or "")
        dto   = parse_date(request.GET.get('to') or "")
        if dfrom:
            qs = qs.filter(created_at__date__gte=dfrom)
        if dto:
            qs = qs.filter(created_at__date__lte=dto)

        q = (request.GET.get('q') or "").strip()
        if q:
            qs = qs.filter(
                Q(header__sale_id__icontains=q) |
                Q(reversed_by__username__icontains=q) |
                Q(reason__icontains=q)
            )

        qs = qs.order_by('-created_at')[:200]  # cap response size

        data = [{
            "sale_id": str(r.header.sale_id),
            "total": float(r.header.total or 0),
            "reversed_by": r.reversed_by.username if r.reversed_by_id else None,
            "reason": r.reason or "",
            "reversed_at": timezone.localtime(r.created_at).strftime("%Y-%m-%d %H:%M"),
        } for r in qs]

        return JsonResponse(data, safe=False, status=200)


# ---------- shared helpers ----------
def _active_branch(request):
    bid = request.session.get("active_branch_id")
    return Branch.objects.filter(pk=bid).first() if bid else None

def _date_range(request, default_days=30):
    """Returns naive dates (YYYY-MM-DD) for [from, to] inclusive."""
    today = timezone.localdate()
    dfrom = parse_date(request.GET.get("from") or "")
    dto   = parse_date(request.GET.get("to") or "")
    if not dfrom and not dto:
        dfrom = today - timezone.timedelta(days=default_days)
        dto = today
    elif dfrom and not dto:
        dto = today
    elif dto and not dfrom:
        dfrom = today - timezone.timedelta(days=default_days)
    return dfrom, dto

def _response_csv(filename: str) -> HttpResponse:
    resp = HttpResponse(content_type="text/csv")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp

def _response_xlsx(filename: str) -> HttpResponse:
    resp = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


# ===============================
# INVENTORY (product-level) — CSV/Excel
# ===============================
class InventoryExportCSV(LoginRequiredMixin, View):
    def get(self, request):
        branch = _active_branch(request)
        q = (request.GET.get("q") or "").strip()

        
        qs = (
            StockLedger.objects
            .values(
                "product__id",
                "product__sku",
                "product__name",
                "product__supplier__name",
                "branch__name",
                "branch_id",
            )
            .annotate(
                on_hand=Coalesce(
                    Sum("qty_change"),
                    Value(0,output_field=DecimalField(max_digits=14,decimal_places=4))
                ),

                cost=Coalesce(
                    Max("batch__buying_cost"),
                    Value(0,output_field=DecimalField(max_digits=14,decimal_places=4))
                ),

                value=Coalesce(
                    Sum(
                        ExpressionWrapper(
                            F("batch__qty_on_hand") *
                            F("batch__buying_cost"),
                            output_field=DecimalField(
                                max_digits=18,
                                decimal_places=2
                            )
                        )
                    ),
                    Value(0,output_field=DecimalField(max_digits=14,decimal_places=4))
                )
            )
            .order_by("product__name")
        )

        if branch:
            qs = qs.filter(branch_id=branch.id)
        if q:
            qs = qs.filter(
                Q(product__name__icontains=q) |
                Q(product__sku__icontains=q) |
                Q(product__supplier__name__icontains=q)
            )

        resp = _response_csv("inventory.csv")
        w = csv.writer(resp)
        w.writerow(["SKU", "Name", "Supplier", "Branch", "Qty On Hand", "Value"])
        for r in qs:
            w.writerow([
                r["product__sku"],
                r["product__name"],
                r["product__supplier__name"] or "",
                r["branch__name"],
                float(r["on_hand"] or 0),
                float(r["value"] or 0),
            ])
        return resp


class InventoryExportExcel(LoginRequiredMixin, View):
    def get(self, request):
        branch = _active_branch(request)
        q = (request.GET.get("q") or "").strip()

        qs = (
            StockLedger.objects
            .values(
                "product__id",
                "product__sku",
                "product__name",
                "product__supplier__name",
                "branch__name",
                "branch_id",
            )
            .annotate(
                on_hand=Coalesce(
                    Sum("qty_change"),
                    Value(0,output_field=DecimalField(max_digits=14,decimal_places=4))
                ),

                cost=Coalesce(
                    Max("batch__buying_cost"),
                    Value(0,output_field=DecimalField(max_digits=14,decimal_places=4))
                ),

                value=Coalesce(
                    Sum(
                        ExpressionWrapper(
                            F("batch__qty_on_hand") *
                            F("batch__buying_cost"),
                            output_field=DecimalField(
                                max_digits=18,
                                decimal_places=2
                            )
                        )
                    ),
                    Value(0,output_field=DecimalField(max_digits=14,decimal_places=4))
                )
            )
            .order_by("product__name")
        )
        if branch:
            qs = qs.filter(branch_id=branch.id)
        if q:
            qs = qs.filter(
                Q(product__name__icontains=q) |
                Q(product__sku__icontains=q) |
                Q(product__supplier__name__icontains=q)
            )

        wb = Workbook()
        ws = wb.active
        ws.title = "Inventory"
        ws.append(["SKU", "Name", "Supplier", "Branch", "Qty On Hand", "Value"])
        for r in qs:
            ws.append([
                r["product__sku"],
                r["product__name"],
                r["product__supplier__name"] or "",
                r["branch__name"],
                float(r["on_hand"] or 0),
                float(r["value"] or 0),
            ])

        resp = _response_xlsx("inventory.xlsx")
        wb.save(resp)
        return resp


# ===============================
# PURCHASES — CSV/Excel (already added earlier; kept here for completeness)
# ===============================
class PurchaseExportCSV(LoginRequiredMixin, View):
    def get(self, request):
        branch = _active_branch(request)
        dfrom, dto = _date_range(request, default_days=90)

        qs = (
            PurchaseHeader.objects
            .select_related("supplier", "branch", "created_by")
            .filter(purchase_date__range=[dfrom, dto])
            .order_by("-created_at")
        )
        if branch:
            qs = qs.filter(branch=branch)

        resp = _response_csv("purchases.csv")
        w = csv.writer(resp)
        w.writerow(["Purchase ID", "Date", "Supplier", "Reference", "Branch", "Total", "Created By"])
        for p in qs:
            w.writerow([
                str(p.purchase_id),
                p.purchase_date.isoformat(),
                p.supplier.name if p.supplier_id else "",
                p.reference or "",
                p.branch.name,
                float(p.total or 0),
                p.created_by.username,
            ])
        return resp


class PurchaseExportExcel(LoginRequiredMixin, View):
    def get(self, request):
        branch = _active_branch(request)
        dfrom, dto = _date_range(request, default_days=90)

        qs = (
            PurchaseHeader.objects
            .select_related("supplier", "branch", "created_by")
            .prefetch_related("lines__product", "lines__batch")
            .filter(purchase_date__range=[dfrom, dto])
            .order_by("-created_at")
        )
        if branch:
            qs = qs.filter(branch=branch)

        wb = Workbook()

        ws1 = wb.active
        ws1.title = "Purchases"
        ws1.append(["Purchase ID", "Date", "Supplier", "Reference", "Branch", "Total", "Created By"])

        for p in qs:
            ws1.append([
                str(p.purchase_id),
                p.purchase_date.isoformat(),
                p.supplier.name if p.supplier_id else "",
                p.reference or "",
                p.branch.name,
                float(p.total or 0),
                p.created_by.username,
            ])

        ws2 = wb.create_sheet("Lines")
        ws2.append(["Purchase ID", "SKU", "Name", "Batch", "Expiry", "Qty", "Cost", "Line Total"])

        for p in qs:
            for l in p.lines.all():
                ws2.append([
                    str(p.purchase_id),
                    l.product.sku,
                    l.product.name,
                    l.batch.batch_no if l.batch_id else "",
                    l.batch.expiry_date.isoformat() if (l.batch_id and l.batch.expiry_date) else "",
                    float(l.qty),
                    float(l.buying_cost),
                    float(l.line_total or (l.qty * l.buying_cost)),
                ])

        resp = _response_xlsx("purchases.xlsx")
        wb.save(resp)
        return resp


# ===============================
# SALES — CSV/Excel
# ===============================
class SalesExportCSV(LoginRequiredMixin, View):
    def get(self, request):
        branch = _active_branch(request)
        dfrom, dto = _date_range(request, default_days=90)
        status = (request.GET.get("status") or "").strip().upper()  # e.g., POSTED

        qs = (
            SalesHeader.objects
            .select_related("branch", "cashier")
            .filter(created_at__date__range=[dfrom, dto])
            .order_by("-created_at")
        )
        if branch:
            qs = qs.filter(branch=branch)
        if status:
            qs = qs.filter(status=status)

        resp = _response_csv("sales.csv")
        w = csv.writer(resp)
        w.writerow(["Sale ID", "DateTime", "Branch", "Cashier", "POS", "Status", "Total", "Tax"])
        for h in qs:
            w.writerow([
                str(h.sale_id),
                timezone.localtime(h.created_at).isoformat(),
                h.branch.name,
                h.cashier.username,
                h.pos_terminal,
                h.status,
                float(h.total or 0),
                float(h.tax_total or 0),
            ])
        return resp


class SalesExportExcel(LoginRequiredMixin, View):
    def get(self, request):
        branch = _active_branch(request)
        dfrom, dto = _date_range(request, default_days=90)
        status = (request.GET.get("status") or "").strip().upper()

        qs = (
            SalesHeader.objects
            .select_related("branch", "cashier")
            .prefetch_related("lines__product", "lines__batch")
            .filter(created_at__date__range=[dfrom, dto])
            .order_by("-created_at")
        )
        if branch:
            qs = qs.filter(branch=branch)
        if status:
            qs = qs.filter(status=status)

        wb = Workbook()

        ws1 = wb.active
        ws1.title = "Sales"
        ws1.append(["Sale ID", "DateTime", "Branch", "Cashier", "POS", "Status", "Total", "Tax"])
        for h in qs:
            ws1.append([
                str(h.sale_id),
                timezone.localtime(h.created_at).strftime("%Y-%m-%d %H:%M:%S"),
                h.branch.name,
                h.cashier.username,
                h.pos_terminal,
                h.status,
                float(h.total or 0),
                float(h.tax_total or 0),
            ])

        ws2 = wb.create_sheet("Lines")
        ws2.append(["Sale ID", "SKU", "Name", "Batch", "Expiry", "Qty", "Unit Price", "Discount", "Line Total"])
        for h in qs:
            for l in h.lines.all():
                ws2.append([
                    str(h.sale_id),
                    l.product.sku,
                    l.product.name,
                    l.batch.batch_no if l.batch_id else "",
                    l.batch.expiry_date.isoformat() if (l.batch_id and l.batch.expiry_date) else "",
                    float(l.qty),
                    float(l.unit_price),
                    float(l.discount or 0),
                    float(l.line_total or (l.unit_price * l.qty - (l.discount or 0))),
                ])

        resp = _response_xlsx("sales.xlsx")
        wb.save(resp)
        return resp


# ===============================
# LEDGER — CSV/Excel
# ===============================
class LedgerExportCSV(LoginRequiredMixin, View):
    def get(self, request):
        branch = _active_branch(request)
        dfrom, dto = _date_range(request, default_days=90)
        reason = (request.GET.get("reason") or "").strip().upper()  # IN/OUT
        q = (request.GET.get("q") or "").strip()

        qs = (
            StockLedger.objects
            .select_related("product", "branch", "batch")
            .filter(created_at__date__range=[dfrom, dto])
            .order_by("-created_at")
        )
        if branch:
            qs = qs.filter(branch=branch)
        if reason in ("IN", "OUT"):
            qs = qs.filter(reason=reason)
        if q:
            qs = qs.filter(
                Q(product__sku__icontains=q) |
                Q(product__name__icontains=q) |
                Q(reference__icontains=q)
            )

        resp = _response_csv("stock_ledger.csv")
        w = csv.writer(resp)
        w.writerow(["DateTime", "Branch", "SKU", "Product", "Batch", "Expiry", "Qty Change", "Reason", "Reference"])
        for x in qs:
            w.writerow([
                timezone.localtime(x.created_at).strftime("%Y-%m-%d %H:%M:%S"),
                x.branch.name,
                x.product.sku,
                x.product.name,
                x.batch.batch_no if x.batch_id else "",
                x.batch.expiry_date.isoformat() if (x.batch_id and x.batch.expiry_date) else "",
                float(x.qty_change),
                x.reason,
                x.reference or "",
            ])
        return resp


class LedgerExportExcel(LoginRequiredMixin, View):
    def get(self, request):
        branch = _active_branch(request)
        dfrom, dto = _date_range(request, default_days=90)
        reason = (request.GET.get("reason") or "").strip().upper()
        q = (request.GET.get("q") or "").strip()

        qs = (
            StockLedger.objects
            .select_related("product", "branch", "batch")
            .filter(created_at__date__range=[dfrom, dto])
            .order_by("-created_at")
        )
        if branch:
            qs = qs.filter(branch=branch)
        if reason in ("IN", "OUT"):
            qs = qs.filter(reason=reason)
        if q:
            qs = qs.filter(
                Q(product__sku__icontains=q) |
                Q(product__name__icontains=q) |
                Q(reference__icontains=q)
            )

        wb = Workbook()
        ws = wb.active
        ws.title = "Stock Ledger"
        ws.append(["DateTime", "Branch", "SKU", "Product", "Batch", "Expiry", "Qty Change", "Reason", "Reference"])
        for x in qs:
            ws.append([
                timezone.localtime(x.created_at).strftime("%Y-%m-%d %H:%M:%S"),
                x.branch.name,
                x.product.sku,
                x.product.name,
                x.batch.batch_no if x.batch_id else "",
                x.batch.expiry_date.isoformat() if (x.batch_id and x.batch.expiry_date) else "",
                float(x.qty_change),
                x.reason,
                x.reference or "",
            ])

        resp = _response_xlsx("stock_ledger.xlsx")
        wb.save(resp)
        return resp


# ===============================
# STOCK MOVEMENT HISTORY — CSV/Excel
# (can be same as ledger but keeping separate route for clarity)
# ===============================
class StockMovementExportCSV(LedgerExportCSV):
    def get(self, request):
        return super().get(request)  # same data / filters as ledger


class StockMovementExportExcel(LedgerExportExcel):
    def get(self, request):
        return super().get(request)  # same data / filters as ledger


# ===============================
# EXPIRY / BATCH INVENTORY — CSV/Excel
# ===============================
class ExpiryBatchExportCSV(LoginRequiredMixin, View):
    def get(self, request):
        branch = _active_branch(request)
        days = int(request.GET.get("days") or 90)  # next X days
        today = timezone.localdate()
        horizon = today + timezone.timedelta(days=days)

        qs = (
            StockBatch.objects
            .select_related("product", "branch")
            .filter(qty_on_hand__gt=0)
            .order_by("expiry_date", "product__name")
        )
        if branch:
            qs = qs.filter(branch_id=branch.id)

        resp = _response_csv("expiry_batches.csv")
        w = csv.writer(resp)
        w.writerow(["SKU", "Product", "Branch", "Batch", "Expiry", "Days to Expiry", "Qty", "Buying Cost", "Value", "Status"])

        for b in qs:
            exp = b.expiry_date
            days_left = (exp - today).days if exp else ""
            status = "OK"
            if exp:
                if exp <= today:
                    status = "EXPIRED"
                elif exp <= horizon:
                    status = "NEAR-EXPIRY"
            value = (b.qty_on_hand or Decimal("0")) * (b.buying_cost or Decimal("0"))

            w.writerow([
                b.product.sku,
                b.product.name,
                b.branch.name,
                b.batch_no,
                exp.isoformat() if exp else "",
                days_left,
                float(b.qty_on_hand),
                float(b.buying_cost or 0),
                float(value),
                status,
            ])
        return resp


class ExpiryBatchExportExcel(LoginRequiredMixin, View):
    def get(self, request):
        branch = _active_branch(request)
        days = int(request.GET.get("days") or 90)
        today = timezone.localdate()
        horizon = today + timezone.timedelta(days=days)

        qs = (
            StockBatch.objects
            .select_related("product", "branch")
            .filter(qty_on_hand__gt=0)
            .order_by("expiry_date", "product__name")
        )
        if branch:
            qs = qs.filter(branch_id=branch.id)

        wb = Workbook()
        ws = wb.active
        ws.title = "Expiry / Batches"
        ws.append(["SKU", "Product", "Branch", "Batch", "Expiry", "Days to Expiry", "Qty", "Buying Cost", "Value", "Status"])

        for b in qs:
            exp = b.expiry_date
            days_left = (exp - timezone.localdate()).days if exp else ""
            status = "OK"
            if exp:
                if exp <= timezone.localdate():
                    status = "EXPIRED"
                elif exp <= horizon:
                    status = "NEAR-EXPIRY"
            value = (b.qty_on_hand or Decimal("0")) * (b.buying_cost or Decimal("0"))

            ws.append([
                b.product.sku,
                b.product.name,
                b.branch.name,
                b.batch_no,
                exp.isoformat() if exp else "",
                days_left,
                float(b.qty_on_hand),
                float(b.buying_cost or 0),
                float(value),
                status,
            ])

        resp = _response_xlsx("expiry_batches.xlsx")
        wb.save(resp)
        return resp