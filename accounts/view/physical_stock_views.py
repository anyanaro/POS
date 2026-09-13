from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from accounts.models.org import Branch
from accounts.models.stock import StockBatch
from accounts.services.stock_adjustment_engine import adjust_stock


def _branch(request):
    branch = getattr(request, "active_branch", None)
    if branch:
        return branch
    return Branch.objects.filter(pk=request.session.get("active_branch_id")).first()


@login_required
def physical_stock_take(request):
    branch = _branch(request)
    if not branch:
        messages.error(request, "Select an active branch before starting stock take.")
        return redirect("accounts:inventory")

    if request.method == "POST":
        upload = request.FILES.get("file")
        if not upload or not upload.name.lower().endswith(".xlsx"):
            messages.error(request, "Upload the completed .xlsx stock-take file.")
            return redirect("accounts:physical-stock-take")
        try:
            workbook = load_workbook(upload, read_only=True, data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            headers = {str(value).strip().lower(): index for index, value in enumerate(rows[0]) if value is not None}
            required = {"batch_id", "product_sku", "batch_no", "system_qty", "physical_qty"}
            missing = required - headers.keys()
            if missing:
                raise ValueError(f"Missing columns: {', '.join(sorted(missing))}")

            changes = []
            errors = []
            for row_number, values in enumerate(rows[1:], start=2):
                def cell(name):
                    index = headers[name]
                    return values[index] if index < len(values) else None

                if all(cell(name) in (None, "") for name in ("batch_id", "physical_qty")):
                    continue
                try:
                    batch_id = int(cell("batch_id"))
                    physical_qty = Decimal(str(cell("physical_qty")))
                except (TypeError, ValueError, InvalidOperation):
                    errors.append(f"Row {row_number}: batch_id and physical_qty must be valid numbers.")
                    continue
                batch = StockBatch.objects.select_related("product").filter(
                    pk=batch_id,
                    branch=branch,
                ).first()
                if not batch:
                    errors.append(f"Row {row_number}: batch {batch_id} is not in active branch {branch}.")
                    continue
                if physical_qty < 0:
                    errors.append(f"Row {row_number}: physical quantity cannot be negative.")
                    continue
                if str(cell("product_sku")) != batch.product.sku or str(cell("batch_no")) != batch.batch_no:
                    errors.append(f"Row {row_number}: batch identity does not match batch {batch_id}.")
                    continue
                variance = physical_qty - batch.qty_on_hand
                if variance:
                    changes.append({
                        "batch_id": batch.id,
                        "product": str(batch.product),
                        "batch_no": batch.batch_no,
                        "system_qty": str(batch.qty_on_hand),
                        "physical_qty": str(physical_qty),
                        "variance": str(variance),
                    })

            if errors:
                raise ValueError("; ".join(errors[:10]))
            request.session["physical_stock_take_preview"] = {
                "branch_id": branch.id,
                "changes": changes,
            }
            request.session.modified = True
            messages.success(request, f"Preview ready: {len(changes)} adjustments require approval.")
        except Exception as exc:
            messages.error(request, f"Stock take was not posted: {exc}")
        return redirect("accounts:physical-stock-take")

    preview = request.session.get("physical_stock_take_preview")
    if preview and preview.get("branch_id") != branch.id:
        preview = None
    batches = StockBatch.objects.filter(branch=branch).select_related("product").order_by("product__name", "batch_no")
    return render(request, "accounts/inventory/physical_stock_take.html", {
        "branch": branch,
        "batches": batches,
        "preview": preview,
    })


@login_required
def physical_stock_take_approve(request):
    if request.method != "POST":
        return redirect("accounts:physical-stock-take")
    branch = _branch(request)
    preview = request.session.get("physical_stock_take_preview")
    if not branch or not preview or preview.get("branch_id") != branch.id:
        messages.error(request, "The stock-take preview is missing or expired.")
        return redirect("accounts:physical-stock-take")

    try:
        with transaction.atomic():
            for change in preview["changes"]:
                batch = StockBatch.objects.select_for_update().select_related("product").get(
                    pk=change["batch_id"],
                    branch=branch,
                )
                variance = Decimal(change["variance"])
                if batch.qty_on_hand != Decimal(change["system_qty"]):
                    raise ValueError(
                        f"{batch.product} changed after preview. Export and upload a fresh count."
                    )
                adjust_stock(
                    product=batch.product,
                    branch=branch,
                    batch=batch,
                    qty_change=variance,
                    reason="PHYS COUNT",
                    user=request.user,
                )
        del request.session["physical_stock_take_preview"]
        request.session.modified = True
        messages.success(request, f"Stock take approved: {len(preview['changes'])} adjustments posted.")
    except Exception as exc:
        messages.error(request, f"Stock take was not posted: {exc}")
    return redirect("accounts:physical-stock-take")


@login_required
def physical_stock_take_cancel(request):
    request.session.pop("physical_stock_take_preview", None)
    request.session.modified = True
    messages.info(request, "Stock-take preview cancelled.")
    return redirect("accounts:physical-stock-take")


@login_required
def physical_stock_take_export(request):
    branch = _branch(request)
    if not branch:
        messages.error(request, "Select an active branch before exporting stock take.")
        return redirect("accounts:inventory")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Physical Stock Take"
    sheet.append(["batch_id", "product_sku", "product_name", "batch_no", "expiry_date", "system_qty", "physical_qty"])
    for batch in StockBatch.objects.filter(branch=branch).select_related("product").order_by("product__name", "batch_no"):
        sheet.append([
            batch.id,
            batch.product.sku,
            batch.product.name,
            batch.batch_no,
            batch.expiry_date.isoformat() if batch.expiry_date else "",
            float(batch.qty_on_hand),
            "",
        ])
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="physical-stock-take-{branch.code}-{timezone.localdate()}.xlsx"'
    workbook.save(response)
    return response
