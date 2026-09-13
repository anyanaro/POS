from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import load_workbook

from accounts.models.product import Product
from accounts.models.purchase import Supplier

REQUIRED_COLUMNS = {"bc_number", "name", "unit_price"}


def _text(value):
    return str(value).strip() if value is not None else ""


def parse_product_workbook(upload):
    workbook = load_workbook(upload, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], ["The workbook is empty."]

    headers = {_text(value).lower(): index for index, value in enumerate(rows[0]) if _text(value)}
    missing = REQUIRED_COLUMNS - headers.keys()
    if missing:
        return [], [f"Missing required columns: {', '.join(sorted(missing))}."]

    parsed, errors = [], []
    seen_bc_numbers = set()
    for row_number, values in enumerate(rows[1:], start=2):
        def value(column, default=""):
            index = headers.get(column)
            return values[index] if index is not None and index < len(values) else default

        name = _text(value("name"))
        bc_number = _text(value("bc_number"))
        if not name:
            errors.append(f"Row {row_number}: name is required.")
            continue
        if not bc_number:
            errors.append(f"Row {row_number}: bc_number is required.")
            continue
        if bc_number in seen_bc_numbers or Product.objects.filter(bc_number=bc_number).exists():
            errors.append(f"Row {row_number}: BC Number {bc_number} already exists.")
            continue
        try:
            unit_price = Decimal(str(value("unit_price")))
            buying_cost = Decimal(str(value("buying_cost", "0") or "0"))
            reorder_level = Decimal(str(value("reorder_level", "0") or "0"))
            reorder_qty = Decimal(str(value("reorder_qty", "0") or "0"))
        except (InvalidOperation, TypeError, ValueError):
            errors.append(f"Row {row_number}: prices and reorder values must be numeric.")
            continue
        if unit_price < 0 or buying_cost < 0 or reorder_level < 0 or reorder_qty < 0:
            errors.append(f"Row {row_number}: numeric values cannot be negative.")
            continue
        supplier_value = _text(value("supplier"))
        supplier = None
        if supplier_value:
            supplier = Supplier.objects.filter(name__iexact=supplier_value).first()
            if supplier is None and supplier_value.isdigit():
                supplier = Supplier.objects.filter(pk=int(supplier_value)).first()
            if not supplier:
                errors.append(f"Row {row_number}: supplier {supplier_value} was not found.")
                continue
        seen_bc_numbers.add(bc_number)
        parsed.append({
            "bc_number": bc_number,
            "name": name,
            "barcode": _text(value("barcode")),
            "buying_cost": buying_cost,
            "unit_price": unit_price,
            "reorder_level": reorder_level,
            "reorder_qty": reorder_qty,
            "supplier": supplier,
        })
    return parsed, errors


def import_products(upload, commit=True):
    rows, errors = parse_product_workbook(upload)
    if errors or not commit:
        return {"created": 0, "errors": errors, "preview": rows}
    with transaction.atomic():
        for row in rows:
            Product.objects.create(**row)
    return {"created": len(rows), "errors": [], "preview": []}
