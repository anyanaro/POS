# accounts/api/suppliers.py
from rest_framework import serializers, viewsets, permissions, filters
from rest_framework.decorators import action
from django.http import HttpResponse
from django.utils import timezone
from accounts.api.pagination import DefaultPageNumberPagination
from django_filters.rest_framework import DjangoFilterBackend
import csv
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from accounts.models import Product
from accounts.models.purchase import Supplier

class SupplierSerializer(serializers.ModelSerializer):
    class Meta:
        model = Supplier
        fields = ["code", "id", "bc_number", "name", "email", "phone", "address", "active"]

class SupplierViewSet(viewsets.ModelViewSet):
    queryset = Supplier.objects.all().order_by("name")
    serializer_class = SupplierSerializer
    permission_classes = [permissions.IsAuthenticated]

    # 🔎 search + ordering + pagination + filter(active)
    filter_backends = [filters.SearchFilter, filters.OrderingFilter, DjangoFilterBackend]
    search_fields = ["code", "bc_number", "name", "email", "phone", "address"]
    ordering_fields = ["name", "active"]
    ordering = ["name"]
    pagination_class = DefaultPageNumberPagination
    filterset_fields = ["active"]

    @action(detail=False, methods=["get"], url_path="export")
    def export_csv(self, request):
        """Export filtered suppliers to CSV."""
        qs = self.filter_queryset(self.get_queryset())
        ts = timezone.now().strftime("%Y%m%d_%H%M%S")
        filename = f"suppliers_{ts}.csv"

        resp = HttpResponse(content_type="text/csv")
        resp["Content-Disposition"] = f'attachment; filename="{filename}"'

        writer = csv.writer(resp)
        writer.writerow(["Code", "BC Number", "Name", "Email", "Phone", "Address", "Active"])
        for s in qs:
            writer.writerow([
                s.code,
            s.bc_number,
                s.name or "",
                s.email or "",
                s.phone or "",
                (s.address or "").replace("\r\n", " ").replace("\n", " "),
                "Yes" if s.active else "No",
            ])
        return resp

    @action(detail=False, methods=["get"], url_path="export-xlsx")
    def export_xlsx(self, request):
        """Export filtered suppliers to Excel (.xlsx)."""
        qs = self.filter_queryset(self.get_queryset())

        wb = Workbook()
        ws = wb.active
        ws.title = "Suppliers"

        headers = ["Code", "BC Number", "Name", "Email", "Phone", "Address", "Active"]
        ws.append(headers)

        # Basic styling
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="4F46E5")  # indigo
        align_center = Alignment(vertical="center")
        thin = Side(style="thin", color="DDDDDD")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for col_idx, _ in enumerate(headers, start=1):
            c = ws.cell(row=1, column=col_idx)
            c.font = header_font
            c.fill = header_fill
            c.alignment = align_center
            c.border = border

        for s in qs:
            ws.append([
                s.code,
                s.bc_number,
                s.name or "",
                s.email or "",
                s.phone or "",
                (s.address or "").replace("\r\n", " ").replace("\n", " "),
                "Yes" if s.active else "No",
            ])

        # Auto-fit-ish columns (best-effort)
        for column_cells in ws.columns:
            max_len = 0
            col = column_cells[0].column_letter
            for cell in column_cells:
                try:
                    val = str(cell.value) if cell.value is not None else ""
                    max_len = max(max_len, len(val))
                except Exception:
                    pass
            ws.column_dimensions[col].width = min(max(12, max_len + 2), 60)

        # Stream to response
        ts = timezone.now().strftime("%Y%m%d_%H%M%S")
        filename = f"suppliers_{ts}.xlsx"
        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)

        resp = HttpResponse(
            bio.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp["Content-Disposition"] = f'attachment; filename="{filename}"'
        return resp